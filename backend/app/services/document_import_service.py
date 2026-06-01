from io import BytesIO
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PyPdfError
from sqlalchemy.orm import Session

from app.models.knowledge import KnowledgeBase
from app.schemas.knowledge import KnowledgeCreate
from app.services.embedding.base import EmbeddingClient
from app.services.knowledge_service import create_knowledge

MARKDOWN_EXTENSIONS = {".md", ".markdown"}
SUPPORTED_EXTENSIONS = {".pdf", ".docx", *MARKDOWN_EXTENSIONS}
CHUNKING_METHODS = {"fixed_size", "paragraph", "full_document", "markdown_heading"}
MARKDOWN_HEADING_PATTERN = re.compile(r"^ {0,3}#{1,6}\s+\S")
MARKDOWN_FENCE_PATTERN = re.compile(r"^ {0,3}(```|~~~)")


def extract_text_from_document(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension in MARKDOWN_EXTENSIONS:
        return _normalize_text(data.decode("utf-8-sig", errors="ignore"))
    if extension == ".docx":
        return _extract_docx_text(data)
    if extension == ".pdf":
        return _extract_pdf_text(data)
    raise ValueError(f"Unsupported document type: {extension or 'unknown'}")


def split_text_into_chunks(text: str, chunk_size: int = 1200, chunk_overlap: int = 150) -> list[str]:
    normalized = _normalize_text(text)
    if not normalized:
        return []
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")
    if chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size")

    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + chunk_size, len(normalized))
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(normalized):
            break
        start = end - chunk_overlap
    return chunks


def split_text_by_method(
    text: str,
    chunking_method: str = "fixed_size",
    chunk_size: int = 1200,
    chunk_overlap: int = 150,
) -> list[str]:
    normalized = _normalize_text(text)
    if not normalized:
        return []
    if chunking_method not in CHUNKING_METHODS:
        raise ValueError(f"Unsupported chunking method: {chunking_method}")
    if chunking_method == "full_document":
        return [normalized]
    if chunking_method == "markdown_heading":
        return split_markdown_by_headings(normalized, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    if chunking_method == "paragraph":
        chunks: list[str] = []
        for paragraph in normalized.split("\n"):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            if len(paragraph) <= chunk_size:
                chunks.append(paragraph)
            else:
                chunks.extend(split_text_into_chunks(paragraph, chunk_size=chunk_size, chunk_overlap=chunk_overlap))
        return chunks
    return split_text_into_chunks(normalized, chunk_size=chunk_size, chunk_overlap=chunk_overlap)


def split_markdown_by_headings(text: str, chunk_size: int = 1200, chunk_overlap: int = 150) -> list[str]:
    normalized = _normalize_text(text)
    if not normalized:
        return []
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")
    if chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size")

    sections: list[list[str]] = []
    current: list[str] = []
    in_fence = False

    for line in normalized.split("\n"):
        if MARKDOWN_FENCE_PATTERN.match(line):
            in_fence = not in_fence

        if not in_fence and MARKDOWN_HEADING_PATTERN.match(line):
            if current:
                sections.append(current)
            current = [line]
            continue

        if current:
            current.append(line)
        else:
            current = [line]

    if current:
        sections.append(current)

    if len(sections) <= 1 and not MARKDOWN_HEADING_PATTERN.match(sections[0][0]):
        return split_text_into_chunks(normalized, chunk_size=chunk_size, chunk_overlap=chunk_overlap)

    chunks: list[str] = []
    for section in sections:
        section_text = "\n".join(section).strip()
        if not section_text:
            continue
        if len(section_text) <= chunk_size:
            chunks.append(section_text)
        else:
            chunks.extend(split_text_into_chunks(section_text, chunk_size=chunk_size, chunk_overlap=chunk_overlap))
    return chunks


def build_knowledge_payloads_from_document(
    filename: str,
    data: bytes,
    category: str,
    source: str,
    document_path: str | None = None,
    chunking_method: str = "fixed_size",
    chunk_size: int = 1200,
    chunk_overlap: int = 150,
) -> list[KnowledgeCreate]:
    text = extract_text_from_document(filename, data)
    chunks = split_text_by_method(
        text,
        chunking_method=chunking_method,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    if not chunks:
        if Path(filename).suffix.lower() == ".pdf":
            raise ValueError(
                "PDF 中未提取到可用于建库的文本。它可能是扫描件或图片版 PDF，请先 OCR 识别为可复制文字的 PDF 后再上传。"
            )
        raise ValueError("文档中没有可导入的文本内容。")

    stem = Path(filename).stem or "上传文档"
    source_with_filename = f"{source}: {filename}" if source else f"上传文档: {filename}"
    total = len(chunks)
    return [
        KnowledgeCreate(
            title=f"{stem}（第 {index}/{total} 段）" if total > 1 else stem,
            category=category,
            content=chunk,
            source=source_with_filename,
            document_name=filename,
            document_path=document_path,
            chunk_index=index,
            chunk_total=total,
            chunking_method=chunking_method,
        )
        for index, chunk in enumerate(chunks, start=1)
    ]


async def import_document_to_knowledge(
    db: Session,
    filename: str,
    data: bytes,
    category: str,
    source: str,
    embedding_client: EmbeddingClient,
    document_path: str | None = None,
    chunking_method: str = "fixed_size",
    chunk_size: int = 1200,
    chunk_overlap: int = 150,
) -> list[KnowledgeBase]:
    payloads = build_knowledge_payloads_from_document(
        filename=filename,
        data=data,
        category=category,
        source=source,
        document_path=document_path,
        chunking_method=chunking_method,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    imported: list[KnowledgeBase] = []
    for payload in payloads:
        imported.append(await create_knowledge(db, payload, embedding_client))
    return imported


def _extract_docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _normalize_text("\n".join(parts))


def _extract_pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
    except PyPdfError as exc:
        raise ValueError("无法读取 PDF 文件，请确认文件未损坏且不是受密码保护的 PDF。") from exc

    if reader.is_encrypted:
        try:
            decrypt_result = reader.decrypt("")
        except PyPdfError as exc:
            raise ValueError("PDF 文件已加密，请先解除密码保护后再上传。") from exc
        if decrypt_result == 0:
            raise ValueError("PDF 文件已加密，请先解除密码保护后再上传。")

    page_text = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except PyPdfError as exc:
            raise ValueError(f"PDF 第 {page_number} 页文本提取失败，请尝试另存为标准 PDF 后重新上传。") from exc
        if text.strip():
            page_text.append(text)
    return _normalize_text("\n".join(page_text))


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return "\n".join(line for line in lines if line)
