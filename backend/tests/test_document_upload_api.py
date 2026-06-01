from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api.routes.auth import get_current_user
from app.core.config import Settings, get_settings
from app.core.database import Base, get_db
from app.main import app
from app.models import Feedback, KnowledgeBase, QARecord, User  # noqa: F401


def test_upload_document_imports_docx_into_knowledge_base(tmp_path, monkeypatch):
    db_path = tmp_path / "test.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    monkeypatch.setenv("EMBEDDING_PROVIDER", "mock")
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = lambda: User(
        id=1,
        username="test-user",
        hashed_password="unused",
        nickname="Test User",
    )
    upload_dir = tmp_path / "uploads"
    app.dependency_overrides[get_settings] = lambda: Settings(
        EMBEDDING_PROVIDER="mock",
        MODEL_PROVIDER="mock",
        UPLOAD_DIR=str(upload_dir),
    )

    document = Document()
    document.add_heading("图书馆借阅规则", level=1)
    document.add_paragraph("学生借阅图书需使用本人校园卡，逾期归还按图书馆规定处理。")
    buffer = BytesIO()
    document.save(buffer)

    client = TestClient(app)
    response = client.post(
        "/api/knowledge/import",
        data={
            "category": "图书馆服务",
            "source": "上传测试",
            "chunking_method": "paragraph",
            "chunk_size": "300",
            "chunk_overlap": "0",
        },
        files={
            "file": (
                "library.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    app.dependency_overrides.clear()

    assert response.status_code == 201
    payload = response.json()
    stored_file_path = Path(payload["stored_file_path"])
    assert stored_file_path.exists()
    assert stored_file_path.read_bytes() == buffer.getvalue()
    assert upload_dir in stored_file_path.parents
    assert payload["imported_count"] == 2
    assert payload["items"][0]["category"] == "图书馆服务"
    assert payload["items"][0]["document_name"] == "library.docx"
    assert payload["items"][0]["document_path"] == str(stored_file_path)
    assert payload["items"][0]["chunk_index"] == 1
    assert payload["items"][0]["chunk_total"] == 2
    assert payload["items"][0]["chunking_method"] == "paragraph"
    assert "图书馆借阅规则" in payload["items"][0]["content"]
    assert "学生借阅图书" in payload["items"][1]["content"]


def test_upload_blank_pdf_returns_actionable_error_and_removes_file(tmp_path, monkeypatch):
    db_path = tmp_path / "test.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    monkeypatch.setenv("EMBEDDING_PROVIDER", "mock")
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = lambda: User(
        id=1,
        username="test-user",
        hashed_password="unused",
        nickname="Test User",
    )
    upload_dir = tmp_path / "uploads"
    app.dependency_overrides[get_settings] = lambda: Settings(
        EMBEDDING_PROVIDER="mock",
        MODEL_PROVIDER="mock",
        UPLOAD_DIR=str(upload_dir),
    )

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = BytesIO()
    writer.write(buffer)

    client = TestClient(app)
    response = client.post(
        "/api/knowledge/import",
        data={
            "category": "综合服务",
            "source": "上传测试",
            "chunking_method": "paragraph",
            "chunk_size": "300",
            "chunk_overlap": "0",
        },
        files={"file": ("scanned.pdf", buffer.getvalue(), "application/pdf")},
    )

    app.dependency_overrides.clear()

    assert response.status_code == 400
    assert "PDF 中未提取到可用于建库的文本" in response.json()["detail"]
    assert [path for path in upload_dir.rglob("*") if path.is_file()] == []
