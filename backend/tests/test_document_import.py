from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.document_import_service import (
    build_knowledge_payloads_from_document,
    extract_text_from_document,
    split_text_by_method,
    split_text_into_chunks,
)


def test_extract_text_from_markdown_bytes():
    content = "# 校园网络办理指南\n\n学生可通过网上办事大厅申请校园网账号。"

    text = extract_text_from_document("network.md", content.encode("utf-8"))

    assert "校园网络办理指南" in text
    assert "网上办事大厅" in text


def test_extract_text_from_docx_bytes():
    document = Document()
    document.add_heading("学生请假流程", level=1)
    document.add_paragraph("学生请假需在辅导员审批后提交学院备案。")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_text_from_document("leave.docx", buffer.getvalue())

    assert "学生请假流程" in text
    assert "辅导员审批" in text


def test_extract_text_from_pdf_bytes(monkeypatch):
    class FakePage:
        def __init__(self, text: str):
            self._text = text

        def extract_text(self):
            return self._text

    class FakePdfReader:
        def __init__(self, _stream):
            self.is_encrypted = False
            self.pages = [FakePage("校医院挂号流程\n可先到服务窗口咨询。")]

    monkeypatch.setattr("app.services.document_import_service.PdfReader", FakePdfReader)

    text = extract_text_from_document("hospital.pdf", b"%PDF fake for dispatch test")

    assert "校医院挂号流程" in text
    assert "服务窗口" in text


def test_split_text_into_chunks_keeps_overlap():
    chunks = split_text_into_chunks("一二三四五六七八九十", chunk_size=5, chunk_overlap=2)

    assert chunks == ["一二三四五", "四五六七八", "七八九十"]


def test_split_text_by_paragraph_groups_paragraphs_without_breaking_lines():
    text = "第一段介绍校园卡。\n第二段介绍挂失。\n第三段介绍补办。"

    chunks = split_text_by_method(text, chunking_method="paragraph", chunk_size=18, chunk_overlap=0)

    assert chunks == ["第一段介绍校园卡。", "第二段介绍挂失。", "第三段介绍补办。"]


def test_split_text_by_full_document_returns_single_chunk():
    text = "第一段。\n第二段。"

    chunks = split_text_by_method(text, chunking_method="full_document", chunk_size=5, chunk_overlap=0)

    assert chunks == ["第一段。\n第二段。"]


def test_split_text_by_markdown_heading_keeps_heading_sections():
    text = "# 校园卡指南\n适用于学生。\n## 挂失流程\n先冻结校园卡。\n## 补办材料\n携带身份证件。"

    chunks = split_text_by_method(text, chunking_method="markdown_heading", chunk_size=200, chunk_overlap=0)

    assert chunks == [
        "# 校园卡指南\n适用于学生。",
        "## 挂失流程\n先冻结校园卡。",
        "## 补办材料\n携带身份证件。",
    ]


def test_split_text_by_markdown_heading_ignores_headings_inside_code_blocks():
    text = "# 网络服务\n```md\n# 这里是代码示例\n```\n正文说明。\n## 账号申请\n通过网上办事大厅申请。"

    chunks = split_text_by_method(text, chunking_method="markdown_heading", chunk_size=200, chunk_overlap=0)

    assert chunks == [
        "# 网络服务\n```md\n# 这里是代码示例\n```\n正文说明。",
        "## 账号申请\n通过网上办事大厅申请。",
    ]


def test_build_knowledge_payloads_from_document_uses_filename_metadata():
    content = "校园卡挂失流程。" * 80

    payloads = build_knowledge_payloads_from_document(
        filename="campus-card.md",
        data=content.encode("utf-8"),
        category="校园卡服务",
        source="上传文档",
        chunking_method="fixed_size",
        chunk_size=120,
        chunk_overlap=20,
    )

    assert len(payloads) >= 2
    assert payloads[0].category == "校园卡服务"
    assert payloads[0].source == "上传文档: campus-card.md"
    assert payloads[0].title.startswith("campus-card")
    assert payloads[0].document_name == "campus-card.md"
    assert payloads[0].chunking_method == "fixed_size"
    assert payloads[0].chunk_index == 1
    assert payloads[0].chunk_total == len(payloads)


def test_build_knowledge_payloads_from_document_supports_paragraph_method():
    content = "校园卡挂失流程。\n图书馆开放时间。\n宿舍报修流程。"

    payloads = build_knowledge_payloads_from_document(
        filename="guide.md",
        data=content.encode("utf-8"),
        category="综合服务",
        source="上传文档",
        chunking_method="paragraph",
        chunk_size=20,
        chunk_overlap=0,
    )

    assert [payload.content for payload in payloads] == [
        "校园卡挂失流程。",
        "图书馆开放时间。",
        "宿舍报修流程。",
    ]
    assert all(payload.document_name == "guide.md" for payload in payloads)
    assert all(payload.chunking_method == "paragraph" for payload in payloads)


def test_extract_text_rejects_unsupported_extension():
    with pytest.raises(ValueError, match="Unsupported document type"):
        extract_text_from_document("notes.txt", b"hello")


def test_build_knowledge_payloads_from_blank_pdf_returns_actionable_error():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(ValueError, match="PDF 中未提取到可用于建库的文本"):
        build_knowledge_payloads_from_document(
            filename="scanned.pdf",
            data=buffer.getvalue(),
            category="综合服务",
            source="上传文档",
        )
